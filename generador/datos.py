"""Carga test, supuestos e infracciones de Local Andalucía y descarta lo que no se debe publicar."""
import glob
import os
import re

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, "..", ".."))

CARPETA_TEST = os.path.join(RAIZ, "10. TEST", "20. SOLUCIÓN EXCEL", "BBDD CORREGIDA")
CARPETA_CORTAS = os.path.join(RAIZ, "11. SUPUESTOS PRÁCTICOS", "20. SOLUCIÓN EXCEL SP", "SUBIDOS")
CARPETA_CODIFICADO = os.path.join(RAIZ, "20. RECURSOS IA", "CODIFICADO")

# Limites para que el texto quepa con un tamano legible
MAX_PREGUNTA = 300
MAX_OPCION = 220
MAX_ENUNCIADO = 1500  # dos diapositivas como maximo
MAX_RESPUESTA = 1400  # respuesta modelo de los supuestos de desarrollo
MAX_INFRACCION = 420

# Normas derogadas o datos que cambian con el tiempo: no publicar sin revisar
DESFASADAS = [
    r"ley 13/2001",                    # coordinación de policías locales de Andalucía -> Ley 6/2023
    r"ley org[aá]nica 1/1992",         # seguridad ciudadana -> LO 4/2015
    r"real decreto legislativo 339/1990", r"rdl 339/1990",  # tráfico -> RDL 6/2015
    r"ley 30/1992", r"ley 11/2007", r"ley 6/1997", r"ley 30/1984",
    r"eurodiputad",
]

NORMAS_CORTAS = {
    "Constitución Española": "CE",
    "Código Penal": "Código Penal",
    "10/1995": "Código Penal",
    "Enjuiciamiento Criminal": "LECrim",
    "Ley 39/2015": "Ley 39/2015",
    "Ley 40/2015": "Ley 40/2015",
    "Ley 7/1985": "LRBRL",
    "Orgánica 2/1986": "LOFCS",
    "Orgánica 4/2015": "LOPSC",
    "Ley 6/2023": "Ley 6/2023 Policías Locales",
    "Real Decreto Legislativo 6/2015": "LSV",
    "Real Decreto 1428/2003": "RGC",
    "Real Decreto 818/2009": "RGCond",
    "Real Decreto 2822/1998": "RGV",
    "Estatuto de Autonomía": "EAA",
    "Real Decreto Legislativo 5/2015": "TREBEP",
    "Orgánica 3/2007": "LO 3/2007",
    "Orgánica 1/2004": "LO 1/2004",
    "Orgánica 5/2000": "LORPM",
    "Orgánica 6/1984": "LO Habeas Corpus",
    "Orgánica 4/2000": "LO Extranjería",
    "Orgánica 3/2018": "LOPDGDD",
    "Ley 7/2023": "Ley 7/2023 Animales",
}

TILDES = {"Cadiz": "Cádiz", "Malaga": "Málaga", "Cordoba": "Córdoba", "Jaen": "Jaén",
          "Diputacion": "Diputación", "Almeria": "Almería", "Gines": "Ginés", "Linea": "Línea",
          "Concepcion": "Concepción", "Velez": "Vélez", "Alcala": "Alcalá", "Chipinoa": "Chipiona",
          "Santo Tome": "Santo Tomé", "Almuñecar": "Almuñécar", "Almodovar": "Almodóvar",
          "Benalmadena": "Benalmádena", "Cartama": "Cártama", "Ubeda": "Úbeda", "Ecija": "Écija",
          "Andujar": "Andújar", "Ogijares": "Ogíjares", "Rio": "Río", "Leon": "León", "Marmolejo": "Marmolejo"}

# Nombres de columna distintos segun la epoca del Excel -> nombre canonico
CANON = {
    "opciona": "Opción A", "opcionb": "Opción B", "opcionc": "Opción C", "opciond": "Opción D",
    "opcioncorrecta": "Respuesta Correcta", "explicacion": "Justificación",
    "justificación correcta": "Justificación",
    "enunciado": "Enunciado", "pregunta corta": "Pregunta", "solución": "Justificación",
    "normativa": "Normativa", "tema": "Tema", "convocatoria": "Convocatoria",
}


INICIO_NORMA = 45  # "Real Decreto 1428/2003, de 21 de noviembre," ; lo que sigue son normas citadas


def derogada(norma):
    return any(re.search(p, norma[:INICIO_NORMA], re.I) for p in DESFASADAS)


def limpiar(t):
    t = "" if t is None else str(t)
    t = t.replace("/n", "\n").strip()
    t = re.sub(r"\*\*|__|^#{1,6}\s*", "", t, flags=re.M)  # restos de Markdown
    t = re.sub(r"^\s*Justificaci[oó]n (correcta|incorrecta)\s*:\s*", "", t, flags=re.I)
    t = re.sub(r"[ \t]+", " ", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def acortar(t, maximo):
    if len(t) <= maximo:
        return t
    corte = t[:maximo].rsplit(" ", 1)[0]
    punto = corte.rfind(". ")
    if punto > maximo * 0.5:
        return corte[:punto + 1]
    return corte + "…"


def tildar(t):
    for k, v in TILDES.items():
        t = re.sub(r"\b" + k + r"\b", v, t)
    return t


def norma_corta(n):
    for k, v in NORMAS_CORTAS.items():
        if k.lower() in n.lower():
            return v
    return n if len(n) <= 30 else ""


def nombre_convocatoria(conv):
    conv = re.sub(r"\s+(PC|SP|TEST|SP TEST|PC TEST)$", "", conv.strip(), flags=re.I)
    m = re.match(r"^(\d{4})\s*-?\s*(.+)$", conv)
    if m:
        conv = f"{m.group(2)} {m.group(1)}"
    conv = re.sub(r"\s*\(.*?\)", "", conv)
    lugar, _, anio = conv.rpartition(" ")
    if anio.isdigit() and len(anio) == 4:
        lugar = re.split(r"\s+(?:primer|segundo|tercer|supuesto|cuestionario|plantilla|promoci|interino|\d)",
                         lugar, maxsplit=1, flags=re.I)[0]
        lugar = re.sub(r"\s+SP$", "", lugar, flags=re.I)
        if lugar.isupper():
            lugar = lugar.title()
        conv = f"{lugar} {anio}"
    menores = {"de", "la", "del", "las", "los", "y", "el"}
    conv = " ".join(w.lower() if (i and w.lower() in menores) else w[:1].upper() + w[1:]
                    for i, w in enumerate(conv.split()))
    return tildar(conv.strip())


def _fila_a_pregunta(r, tipo, origen):
    articulo = r.get("Artículo num.", "")
    if articulo and not articulo.lower().startswith("art"):
        articulo = "Art. " + articulo
    articulo = articulo.replace("Artículo", "Art.")
    norma = r.get("Normativa", "")
    conv = r.get("Convocatoria", "") or os.path.splitext(os.path.basename(origen.split("#")[0]))[0]
    return {
        "id": origen,
        "tipo": tipo,  # "test" | "corta"
        "enunciado": r.get("Enunciado", ""),
        "pregunta": re.sub(r"^RESERVA\s*\d*\s*[.:-]?\s*", "", r.get("Pregunta", ""), flags=re.I),
        "opciones": [r.get(f"Opción {l}", "") for l in "ABCD"],
        "correcta": r.get("Respuesta Correcta", "").strip().upper()[:1],
        "justificacion": r.get("Justificación", ""),
        "justificacion_incorrecta": r.get("Justificación Incorrecta", ""),
        "literal": r.get("Artículo Literal", ""),
        "articulo": articulo,
        "norma": norma,
        "norma_corta": norma_corta(norma) if norma else "",
        "tema": r.get("Tema", ""),
        "convocatoria": nombre_convocatoria(conv),
        "anio": int(m.group(1)) if (m := re.search(r"(20\d\d)", conv)) else 0,
    }


def _cabecera(fila):
    return [CANON.get(limpiar(c).lower(), limpiar(c)) for c in fila]


def _leer_excel(ruta, tipo):
    ws = openpyxl.load_workbook(ruta, read_only=True).worksheets[0]
    filas = ws.iter_rows(values_only=True)
    cab = _cabecera(next(filas))
    base = os.path.relpath(ruta, RAIZ).replace("\\", "/")
    for n, val in enumerate(filas, start=2):
        r = {c: limpiar(v) for c, v in zip(cab, val)}
        if r.get("Pregunta"):
            yield _fila_a_pregunta(r, tipo, f"{base}#{n}")


def es_desarrollo(p):
    return p["tipo"] == "corta" and not any(p["opciones"])


def motivo_pregunta_mala(t):
    """Por que una pregunta de supuesto no se entiende sola en Instagram ('' si se entiende)."""
    t = re.sub(r"^\W*(\d+\s*[ºª]?|[A-H])\s*[.)-]*\s*-?\s*", "", t.strip())
    if re.search(r"\b(imagen|im[aá]genes|fotograf[ií]a|foto|figura|croquis|plano adjunto|gr[aá]fico)\b", t, re.I):
        return "la pregunta depende de una imagen"
    if t.endswith(":"):
        return "la pregunta está incompleta"
    if "?" not in t and len(t) < 45:
        return "la pregunta es un encabezado, no una pregunta"
    return ""


# ---------------------------------------------------------------- infracciones (codificado)

def _texto_infraccion(t):
    # quita las instrucciones al agente: "(DEBERÁ INDICARSE...)", "(DESCRIBIR...)"
    t = re.sub(r"\((?:[^()]*?\b(?:DEBER[AÁ]|DESCRIBIR|INDICAR|ESPECIFICAR|HACER CONSTAR|CONSIGNAR)\b[^()]*)\)", "", t, flags=re.I)
    return re.sub(r"\s+", " ", t).strip(" .") + "."


def resumir(t, maximo):
    """Primera frase util de un campo largo, sin las citas '(art. 62.1.a)'.
    No corta en abreviaturas como 'Seg.' o 'art.': solo tras palabras de 4+ letras."""
    t = re.sub(r"\s*\((?:arts?|apdo|apartado)\.?[^)]*\)", "", t, flags=re.I)
    t = re.split(r"(?<=\w{4}\.)\s+(?=[A-ZÁÉÍÓÚ])", t, maxsplit=1)[0].strip(" ;")
    if len(t) <= maximo:
        return t
    # cierre de frase real antes del limite: "palabra." o ")."
    finales = [m.end() for m in re.finditer(r"(\w{4}|\))\.(?=\s)", t) if m.end() <= maximo]
    return t[:finales[-1]] if finales else acortar(t, maximo)


SIGLAS_NORMA = [("Real Decreto Legislativo", "RDL"), ("Real Decreto-ley", "RDL"), ("Real Decreto", "RD"),
                ("Ley Orgánica", "LO"), ("Decreto Legislativo", "DL")]


def norma_media(n):
    """Nombre legible y corto: 'Reglamento de artículos pirotécnicos y cartuchería (RD 989/2015)'."""
    n = n.strip(" .")
    m = re.match(r"^(.*?\d+/\d{4}),.*?por (?:el|la) que se aprueba (?:el|la)?\s*(.+)$", n, re.I)
    if m:
        ident = m.group(1)
        for largo, sigla in SIGLAS_NORMA:
            ident = ident.replace(largo, sigla)
        titulo = m.group(2)
        return acortar(titulo[:1].upper() + titulo[1:], 80) + f" ({ident})"
    m = re.match(r"^(.*?\d+/\d{4}), de \d+ de \w+,?\s*(.+)$", n)
    if m:
        return acortar(f"{m.group(1)}, {m.group(2)}", 95)
    return acortar(n, 95)


def primera_idea(t, maximo=300):
    """Primera frase del comentario, saltando lo generico (prescripcion, graduacion...)."""
    frases = [f.strip() for f in re.split(r"(?<=[.])\s+(?=[A-ZÁÉÍÓÚ])", t) if f.strip()]
    utiles = [f for f in frases if not re.match(r"(Prescripci|Criterios de graduaci|Graduaci|Procedimiento|Caducidad|Competencia|El plazo|Los plazos|Reincidencia)", f)]
    # solo frases completas: si la primera no cabe entera, mejor no poner comentario
    return utiles[0] if utiles and len(utiles[0]) <= maximo else ""


MATERIAS = {"Mediambiente": "Medioambiente"}


def cargar_infracciones():
    for ruta in sorted(glob.glob(os.path.join(CARPETA_CODIFICADO, "*.xlsx"))):
        if os.path.basename(ruta).startswith("~$"):
            continue
        ws = openpyxl.load_workbook(ruta, read_only=True).worksheets[0]
        filas = ws.iter_rows(values_only=True)
        cab = [limpiar(c) for c in next(filas)]
        base = os.path.relpath(ruta, RAIZ).replace("\\", "/")
        for n, val in enumerate(filas, start=2):
            r = {c: limpiar(v) for c, v in zip(cab, val)}
            if not r.get("Texto"):
                continue
            util = lambda k: r.get(k, "") not in ("", "-", "None")
            ref = " ".join(x for x in (
                f"Art. {r['Artículo']}" if util("Artículo") else "",
                f"apdo. {r['Apartado']}" if util("Apartado") else "",
                f"opc. {r['Subapartado']}" if util("Subapartado") else "") if x)
            puntos = r.get("Puntos", "")
            yield {
                "id": f"{base}#{n}",
                "tipo": "infraccion",
                "materia": MATERIAS.get(r.get("Tema", ""), r.get("Tema", "")) or os.path.basename(ruta).split(" - ")[0].title(),
                "tema": r.get("Tema", ""),
                "norma": r.get("Norma", ""),
                "norma_corta": norma_corta(r.get("Norma", "")),
                "norma_media": norma_corta(r.get("Norma", "")) or norma_media(r.get("Norma", "")),
                "gravedad": r.get("Tipo de infracción", "").strip().capitalize(),
                "referencia": ref,
                "texto": _texto_infraccion(r["Texto"]),
                "cuantia": r.get("Cuantía", ""),
                "puntos": puntos if puntos not in ("0", "0.0", "None") else "",
                "responsable": resumir(r.get("Responsable", ""), 110),
                "organo": resumir(r.get("Órgano Competente", ""), 110),
                "comentario": primera_idea(r.get("Comentario", "")),
                "convocatoria": "", "anio": 0,
            }


def problemas_infraccion(i):
    mal = []
    if i["gravedad"].lower() not in ("leve", "grave", "muy grave"):
        mal.append("gravedad no reconocida")
    if not i["cuantia"]:
        mal.append("sin cuantía")
    if len(i["texto"]) < 25:
        mal.append("texto demasiado corto")
    if len(i["texto"]) > MAX_INFRACCION:
        mal.append("texto demasiado largo para el carrusel")
    if derogada(i["norma"]):
        mal.append("norma derogada")
    if re.search(r"\b(leves?|graves?)\b", i["texto"], re.I):
        mal.append("el texto ya dice la gravedad")
    return mal


# ---------------------------------------------------------------- filtro de calidad

def problemas(p):
    """Lista de motivos por los que un elemento NO debe publicarse (vacia = OK)."""
    if p["tipo"] == "infraccion":
        return problemas_infraccion(p)
    mal = []
    ops = [o for o in p["opciones"] if o]
    if es_desarrollo(p):
        if motivo := motivo_pregunta_mala(p["pregunta"]):
            mal.append(motivo)
        if len(p["justificacion"]) < 80:
            mal.append("respuesta modelo vacía o insuficiente")
        if len(p["justificacion"]) > MAX_RESPUESTA:
            mal.append("respuesta modelo demasiado larga para el carrusel")
    else:
        if len(ops) < 3:
            mal.append("menos de 3 opciones")
        if p["correcta"] not in "ABCD" or not p["correcta"]:
            mal.append("sin respuesta correcta")
        elif not p["opciones"]["ABCD".index(p["correcta"])]:
            mal.append("la respuesta correcta apunta a una opción vacía")
        if not p["justificacion"]:
            mal.append("sin justificación")
        if motivo_pregunta_mala(p["pregunta"]) == "la pregunta depende de una imagen":
            mal.append("la pregunta depende de una imagen")
    todo = " ".join([p["enunciado"], p["pregunta"], *p["opciones"], p["norma"][:INICIO_NORMA]])
    if re.search(r"opci[oó]n omitida|\[.*?(ilegible|omitid).*?\]|\?\?", todo, re.I):
        mal.append("texto incompleto en el original")
    if re.search(r"\b(20[3-9]\d|2[1-9]\d\d)\b(?!/)", " ".join([p["pregunta"], *p["opciones"]])):
        mal.append("año sospechoso (posible errata)")
    for patron in DESFASADAS:
        if re.search(patron, todo, re.I):
            mal.append(f"dato que puede estar desfasado ({patron})")
            break
    if len(p["pregunta"]) > MAX_PREGUNTA:
        mal.append("pregunta demasiado larga para el carrusel")
    if any(len(o) > MAX_OPCION for o in ops):
        mal.append("opción demasiado larga para el carrusel")
    if p["tipo"] == "corta" and len(p["enunciado"]) > MAX_ENUNCIADO:
        mal.append("enunciado demasiado largo para el carrusel")
    if p["tipo"] == "corta" and len(p["enunciado"]) < 60:
        mal.append("enunciado vacío o insuficiente")
    return mal


def cargar(tipo):
    if tipo == "infraccion":
        yield from cargar_infracciones()
        return
    carpeta = CARPETA_TEST if tipo == "test" else CARPETA_CORTAS
    for ruta in sorted(glob.glob(os.path.join(carpeta, "**", "*.xlsx"), recursive=True)):
        if os.path.basename(ruta).startswith("~$"):
            continue
        yield from _leer_excel(ruta, tipo)


def informe():
    from collections import Counter
    for tipo in ("test", "corta", "infraccion"):
        total, ok, motivos = 0, 0, Counter()
        for p in cargar(tipo):
            total += 1
            m = problemas(p)
            if m:
                motivos[m[0].split(" (")[0]] += 1
            else:
                ok += 1
        print(f"\n{tipo.upper()}: {ok} publicables de {total}")
        for k, v in motivos.most_common():
            print(f"   descartadas por {k}: {v}")


if __name__ == "__main__":
    informe()
